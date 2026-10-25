"""Report sources from species other than mouse and human (decision 9.94, the owner, 24 October 2026: "Pooled sources
until standalone").

Each source here is class animal, a soft prior through --soft-prior-reports only and never a label (9.48, 9.50), with
a source and a species of its own (evidence/calibrated_reliability.py SPECIES_OF_SOURCE). The calibrated regression
reads a species through the pooled form with the human cells while its usable calibration cells hold fewer than 20
positives or 20 negatives, and fits it on its own once they hold both (the separate form of 9.93), so no source here
needs a rule of its own for when it stands alone.

Genes. A gene of another species reaches a human gene through the Alliance of Genome Resources 9.0.0 orthology file
(ORTHOLOGY-ALLIANCE_COMBINED.tsv.gz, "Orthology Filter: Stringent" in its header), one to one in both directions: the
species gene has one human orthologue in the file and that human gene has one orthologue of the species
(read_one_to_one_orthologs). The human gene then needs a graph node, as for IMPC. "Most Alliance data is available
under CC0 1.0 Universal (Public Domain Dedication)." (registry.opendata.aws/alliance-genome-resources, read 24 October
2026.)

The sources (one builder each; the rules and why are in decision 9.94):
  - RGD-rat (build_rgd_rat): RGD's rat gene to MP annotations (rattus_genes_mp, a GAF), mapped through IMPC's crosswalk
    and the descendant rule of evidence/mgi_mouse_phenotypes.py. IEA rows are left out; rows qualified by a treatment,
    an onset, a severity or another condition are left out; a NOT row is an absence. Allele direction from the allele
    symbol (rat_allele_direction).
  - ChEMBL-rat-in-vivo (build_chembl_rat): drug-produced catalepsy in rats, read as parkinsonism (moderate: predictive
    validity for the extrapyramidal effects of antipsychotics). An ED50, MED or minimum dose with "=" or "<" is a
    call; a ">" value with its limit, or an inactive comment, is an absence; assays of another drug's catalepsy
    (reversal, antagonism, potentiation) are not read.
  - FlyBase-classical and FlyBase-RNAi (build_flybase): FlyBase's genotype to phenotype file; a classical genotype holds
    alleles of one Drosophila gene only; an RNAi genotype holds one RNAi construct of one gene with a whole-body or
    pan-neuronal driver of configs/flybase_drivers.csv.
  - ZFIN-mutant and ZFIN-morpholino (build_zfin): ZFIN's "Fish with Abnormal and Normal Phenotypes", whose normal rows
    are absences, for fish of one gene with no transgenic construct, in the standard or generic-control environment.
  - WormBase-variation and WormBase-RNAi (build_wormbase): WormBase WS298 phenotype associations, NOT rows included as
    absences. Every worm mapping is of the validity class weak_worm, off by default.
  - Fenckova-2019-RNAi (build_fenckova): the Drosophila light-off jump habituation screen of 286 human intellectual
    disability genes with pan-neuronal RNAi (Fenckova et al. 2019, Biol Psychiatry 86:294, doi 10.1016/j.biopsych.2019.04.029, PMID 31272685; Table S2 and S3).
The term mappings of the fly, fish, worm and habituation sources are in configs/species_phenotype_crosswalk.csv, a
validity class per row in the IMPC crosswalk's style.
"""
from __future__ import annotations

import csv
import gzip
import html
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from mechanistic_pathway_learning.evidence.evidence_reports import REPORT_COLUMNS
from mechanistic_pathway_learning.evidence.impc_knockout_phenotypes import VALIDITY_READ_BY_DEFAULT, VALIDITY_RUBRIC, CrosswalkRow
from mechanistic_pathway_learning.evidence.mgi_mouse_phenotypes import crosswalk_term_map

TAXON = {"rat": "NCBITaxon:10116", "fly": "NCBITaxon:7227", "zebrafish": "NCBITaxon:7955", "worm": "NCBITaxon:6239"}
HUMAN_TAXON = "NCBITaxon:9606"
EVIDENCE_CLASS = "animal"
DEFAULT_CROSSWALK_PATH = Path("configs/species_phenotype_crosswalk.csv")
DEFAULT_DRIVERS_PATH = Path("configs/flybase_drivers.csv")
# validity classes off by default beside weak: a construct read in a species where the readout is far from it
SEPARATE_VALIDITY_CLASSES = ("weak_intellectual_disability", "weak_worm")
EXTRA_COLUMNS = ("rubric_cross_species", "rubric_proxy_validity", "allele_direction", "species", "species_gene", "validity")
SOURCE_SPECIES = {"RGD-rat": "rat", "ChEMBL-rat-in-vivo": "rat", "FlyBase-classical": "fly", "FlyBase-RNAi": "fly", "Fenckova-2019-RNAi": "fly",
                  "ZFIN-mutant": "zebrafish", "ZFIN-morpholino": "zebrafish", "WormBase-variation": "worm", "WormBase-RNAi": "worm"}
RGD_QUALIFIERS_READ = ("", "induces")
ZFIN_ENVIRONMENTS_READ = ("ZDB-EXP-041102-1", "ZDB-EXP-070511-5")  # the standard and the generic-control environments
FLY_RNAI_TAGS = re.compile(r"\[(?:RNAi|GD|KK|HMS|HMC|HMJ|JF|GL|GLV|GLC|TRiP|VSH|shRNA)")
CATALEPSY_OF_ANOTHER_DRUG = re.compile(r"(revers|antagon|inhibit|block|attenuat|prevent|potentiat|protect|reduc).{0,80}catalep|induced catalep|catalepsy induced by", re.I)
CATALEPSY_CALL_TYPES = ("ED50", "MED", "EDmin", "ED")
INACTIVE_COMMENT = re.compile(r"^(inactive|not active|no activity)", re.I)


@dataclass(frozen=True)
class SpeciesCrosswalkRow:
    family: str
    term_ids: tuple[str, ...]
    qualities: tuple[str, ...]  # empty means any
    value: int  # 1 a call, 0 an absence
    target_symptom: str
    relation: str
    validity: str
    note: str


def read_species_crosswalk(path: Path = DEFAULT_CROSSWALK_PATH) -> list[SpeciesCrosswalkRow]:
    rows = []
    with open(path, newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            rows.append(SpeciesCrosswalkRow(row["family"].strip(), tuple(term.strip() for term in row["term_ids"].split(";") if term.strip()),
                                            tuple(quality.strip() for quality in row["qualities"].split(";") if quality.strip()), int(row["value"]),
                                            row["target_symptom"].strip(), row["relation"].strip(), row["validity"].strip(), row.get("note", "")))
    return rows


def matching(crosswalk: list[SpeciesCrosswalkRow], family: str, term: str, quality: str, value: int, validities: tuple[str, ...]) -> list[SpeciesCrosswalkRow]:
    return [row for row in crosswalk if row.family == family and row.value == value and term in row.term_ids
            and (not row.qualities or quality in row.qualities) and row.validity in validities]


def read_one_to_one_orthologs(path: Path, species: str) -> dict[str, tuple[str, str, str]]:
    """species gene id -> (species gene symbol, human HGNC id, human symbol), for the pairs that are one to one in both
    directions among the file's (stringent) rows of that species and human."""
    taxon = TAXON[species]
    pairs = set()
    symbols: dict[str, str] = {}
    with gzip.open(path, "rt", encoding="utf-8") as handle:
        rows = (line.rstrip("\n").split("\t") for line in handle if not line.startswith("#"))
        header = next(rows)
        index = {name: position for position, name in enumerate(header)}
        for fields in rows:
            first, second = (fields[index["Gene1ID"]], fields[index["Gene1Symbol"]], fields[index["Gene1SpeciesTaxonID"]]), \
                            (fields[index["Gene2ID"]], fields[index["Gene2Symbol"]], fields[index["Gene2SpeciesTaxonID"]])
            for human, other in ((first, second), (second, first)):
                if human[2] == HUMAN_TAXON and other[2] == taxon:
                    pairs.add((other[0], human[0]))
                    symbols[other[0]], symbols[human[0]] = other[1], human[1]
    human_of, other_of = defaultdict(set), defaultdict(set)
    for other, human in pairs:
        human_of[other].add(human)
        other_of[human].add(other)
    return {other: (symbols[other], next(iter(humans)), symbols[next(iter(humans))]) for other, humans in human_of.items()
            if len(humans) == 1 and len(other_of[next(iter(humans))]) == 1}


def by_symbol(orthologs: dict[str, tuple[str, str, str]]) -> dict[str, str]:
    """species gene symbol -> human symbol; a symbol held by two species genes is left out."""
    found: dict[str, set[str]] = defaultdict(set)
    for _, (symbol, _, human) in orthologs.items():
        found[symbol].add(human)
    return {symbol: next(iter(humans)) for symbol, humans in found.items() if len(humans) == 1}


def report_row(source: str, perturbation_id: str, perturbation_type: str, nodes: list, symptom: str, relation: str, validity: str, value: int,
               species_gene: str, direction: str) -> dict:
    row = {column: "" for column in REPORT_COLUMNS}
    row.update({
        "perturbation_id": perturbation_id, "perturbation_type": perturbation_type, "perturbation_label": perturbation_id, "symptom": symptom, "relation": relation,
        "evidence_class": EVIDENCE_CLASS, "source": source, "report_value": value, "frequency": None, "frequency_denominator": None, "placebo_flag": False,
        "onset": None, "sex": None, "pubmed_reference_count": 0, "rubric_frequency_known": 0.0, "rubric_evidence_code_pcs": 0.0, "rubric_evidence_code_tas": 0.0,
        "rubric_evidence_code_iea": 0.0, "rubric_placebo_controlled": 0.0, "rubric_curated_synopsis": 0.0, "rubric_log_sample_size": 0.0, "association_type": "",
        "rubric_causal_association": 0.0, "perturbation_nodes": json.dumps(nodes), "evidence_date_source": "release",
        "rubric_cross_species": 1.0, "rubric_proxy_validity": VALIDITY_RUBRIC.get(validity, 0.3 if validity.startswith("weak") else 0.0),
        "allele_direction": direction, "species": SOURCE_SPECIES[source], "species_gene": species_gene, "validity": validity,
    })
    return row


def gene_row(source, human, node, direction, symptom, relation, validity, value, species_gene) -> dict:
    sign = 1.0 if direction == "gain_of_function" else -1.0
    perturbation_id = f"{human}|gain_of_function" if direction == "gain_of_function" else human
    return report_row(source, perturbation_id, "gene", [[node, sign, 1.0]], symptom, relation, validity, value, species_gene, direction)


# --- rat: RGD -----------------------------------------------------------------------------------------------------


def plain_symbol(symbol: str) -> str:
    """An RGD allele symbol without its markup: Fmr1<sup>em2Mcwi</sup> -> Fmr1<em2Mcwi>."""
    text = html.unescape(symbol)
    text = re.sub(r"</?i>", "", text)
    return re.sub(r"<sup>(.*?)</sup>", r"<\1>", text)


def rat_allele_direction(symbol: str, name: str) -> tuple[str | None, str]:
    """(loss_of_function, untyped or None; why), from an RGD allele symbol and its name. A transgene (Tg) is not read, as
    for MGI; a name that states a knockout or a null is loss of function; an endonuclease-mediated (em), targeted (tm),
    ENU or spontaneous (m) allele that states no direction is untyped and keyed to the gene (9.54); a bare gene symbol
    (an annotation to the gene, not an allele) is untyped."""
    plain = plain_symbol(symbol)
    superscript = re.search(r"<(.*)>", plain)
    if "Tg(" in plain or (superscript and superscript.group(1).startswith("Tg")):
        return None, "transgene"
    if re.search(r"knock-?out|null|deletion|frameshift", name, re.I):
        return "loss_of_function", "named a knockout, null or deletion"
    if re.search(r"overexpress|gain of function|constitutively active", name, re.I):
        return None, "a gain of function stated in the name, which the allele symbol does not confirm"
    return "untyped", "no direction stated"


def build_rgd_rat(gaf_path: Path, orthologs: dict, node_of_symbol: dict[str, str], crosswalk: list[CrosswalkRow], children: dict, exclusions: dict,
                  term_names: dict, validities: tuple[str, ...] = VALIDITY_READ_BY_DEFAULT) -> tuple[list[dict], dict]:
    term_map, _ = crosswalk_term_map(crosswalk, validities, children, exclusions)
    human_of_symbol = by_symbol(orthologs)
    counts: dict = defaultdict(int)
    rows = []
    with open(gaf_path, encoding="utf-8") as handle:
        for line in handle:
            if line.startswith("!") or not line.strip():
                continue
            fields = line.rstrip("\n").split("\t")
            counts["rows"] += 1
            symbol, qualifier, term, reference, evidence, name = fields[2], fields[3], fields[4], fields[5], fields[6], fields[9]
            if evidence == "IEA":
                counts["not_read:IEA"] += 1
                continue
            negated = qualifier == "NOT"
            if not negated and qualifier not in RGD_QUALIFIERS_READ:
                counts[f"not_read:qualifier {qualifier}"] += 1
                continue
            mapped = term_map.get(term, [])
            if not mapped:
                continue
            gene = re.split(r"<", plain_symbol(symbol))[0]
            direction, why = rat_allele_direction(symbol, name)
            if direction is None:
                counts[f"not_read:{why}"] += 1
                continue
            human = human_of_symbol.get(gene)
            if human is None or human not in node_of_symbol:
                counts["no_one_to_one_human_gene_in_the_graph"] += 1
                continue
            pubmed = next((part[5:] for part in reference.split("|") if part.startswith("PMID:")), "")
            for crosswalk_row, crosswalk_term in mapped:
                if negated and not crosswalk_row.negative_procedure_codes:
                    counts["not_read:NOT on a row without a measured absence"] += 1
                    continue
                relation = "induces" if negated else crosswalk_row.relation
                row = gene_row("RGD-rat", human, node_of_symbol[human], direction, crosswalk_row.symptom, relation, crosswalk_row.validity, 0 if negated else 1, f"RGD:{fields[1]}")
                row.update({"report_id": f"RGD-rat|{human}|{fields[1]}|{term}|{crosswalk_row.symptom}|{'absent' if negated else crosswalk_row.relation}|{pubmed}",
                            "source_record_id": f"RGD:{fields[1]}", "source_record_label": plain_symbol(symbol), "source_term_id": term,
                            "source_term_label": term_names.get(term, ""), "evidence_code": f"rgd_{evidence.lower()}{'_not' if negated else ''}",
                            "references": f"PMID:{pubmed}" if pubmed else reference, "pubmed_reference_count": 1 if pubmed else 0, "evidence_date": fields[13] if len(fields) > 13 else "",
                            "model_description": f"rat {plain_symbol(symbol)} ({name}); {term} {term_names.get(term, '')}" + ("" if term == crosswalk_term else f" read through {crosswalk_term}") + f"; human orthologue {human} by Alliance one to one",
                            "limitations": f"cross-species: a rat {term_names.get(term, term)} read as {crosswalk_row.symptom} with {crosswalk_row.validity} validity ({crosswalk_row.note})"})
                rows.append(row)
                counts["absences" if negated else "calls"] += 1
    return rows, dict(counts)


# --- rat drugs: ChEMBL in vivo ------------------------------------------------------------------------------------


def build_chembl_rat(assays: list[dict], activities: list[dict], perturbations_of_chembl_id: dict[str, list[str]], validity: str = "moderate") -> tuple[list[dict], dict]:
    """Catalepsy produced by a drug in rats, as parkinsonism (relation induces)."""
    description = {assay["assay_chembl_id"]: assay["description"] or "" for assay in assays}
    organism = {assay["assay_chembl_id"]: assay.get("assay_organism") for assay in assays}
    counts: dict = defaultdict(int)
    rows = []
    for activity in activities:
        counts["activities"] += 1
        assay = activity["assay_chembl_id"]
        text = description.get(assay, "")
        if organism.get(assay) != "Rattus norvegicus" or "catalep" not in text.lower():
            counts["not_read:not a rat catalepsy assay"] += 1
            continue
        if CATALEPSY_OF_ANOTHER_DRUG.search(text):
            counts["not_read:catalepsy produced by another drug"] += 1
            continue
        relation, kind, comment = activity.get("standard_relation"), activity.get("standard_type"), activity.get("activity_comment") or ""
        value = activity.get("standard_value")
        if kind in CATALEPSY_CALL_TYPES and relation in ("=", "<", "<=") and value is not None:
            report_value = 1
        elif (kind in CATALEPSY_CALL_TYPES and relation in (">", ">=") and value is not None) or INACTIVE_COMMENT.search(comment):
            report_value = 0
        else:
            counts["not_read:neither a dose producing catalepsy nor a stated absence"] += 1
            continue
        perturbations = sorted({identifier for chembl_id in (activity.get("parent_molecule_chembl_id"), activity.get("molecule_chembl_id"))
                                for identifier in perturbations_of_chembl_id.get(chembl_id or "", [])})
        if not perturbations:
            counts["not_read:molecule not a perturbation of the tables"] += 1
            continue
        for perturbation_id in perturbations:
            row = report_row("ChEMBL-rat-in-vivo", perturbation_id, "drug", [], "parkinsonism", "induces", validity, report_value, activity.get("molecule_chembl_id") or "", "drug")
            row.update({"report_id": f"ChEMBL-rat-in-vivo|{perturbation_id}|{activity['activity_id']}", "source_record_id": str(activity["activity_id"]),
                        "source_record_label": activity.get("molecule_pref_name") or activity.get("molecule_chembl_id"), "source_term_id": assay,
                        "source_term_label": text[:200], "evidence_code": "chembl_rat_catalepsy" if report_value else "chembl_rat_no_catalepsy",
                        "references": f"ChEMBL {activity.get('document_chembl_id', '')}", "evidence_date": str(activity.get("document_year") or ""),
                        "model_description": f"rat in vivo: {kind} {relation} {value} {activity.get('standard_units') or ''} {comment}".strip(),
                        "limitations": "cross-species: drug-produced catalepsy in rats read as parkinsonism with moderate validity (predictive validity for the extrapyramidal effects of antipsychotics); an absence is at the highest dose tested"})
            rows.append(row)
            counts["absences" if report_value == 0 else "calls"] += 1
    return rows, dict(counts)


# --- fly: FlyBase ---------------------------------------------------------------------------------------------------


def read_drivers(path: Path = DEFAULT_DRIVERS_PATH) -> list[tuple[re.Pattern, str]]:
    with open(path, newline="", encoding="utf-8") as handle:
        return [(re.compile(row["superscript_pattern"]), row["scope"]) for row in csv.DictReader(handle)]


def fly_genotype(symbols: list[str], drivers: list[tuple[re.Pattern, str]]) -> tuple[str | None, str, str]:
    """(the fly gene, FlyBase-classical or FlyBase-RNAi, or the reason it is not read) of a genotype's allele symbols."""
    tools = [symbol for symbol in symbols if symbol.startswith(("Scer\\GAL4[", "Scer\\GAL80[", "Dcr-2[UAS"))]
    rest = [symbol for symbol in symbols if symbol not in tools]
    if any("\\" in symbol for symbol in rest):
        return None, "", "a construct of another species (a silencer, an activator or a human gene)"
    genes = {symbol.split("[")[0] for symbol in rest}
    if len(genes) != 1 or not rest:
        return None, "", "not one Drosophila gene"
    rnai = [symbol for symbol in rest if FLY_RNAI_TAGS.search(symbol)]
    if rnai:
        if len(rnai) != len(rest):
            return None, "", "RNAi beside another allele"
        scopes = [scope for symbol in tools if symbol.startswith("Scer\\GAL4[") for pattern, scope in drivers if pattern.search(symbol.split("[", 1)[1])]
        drivers_present = [symbol for symbol in tools if symbol.startswith("Scer\\GAL4[")]
        if len(drivers_present) != 1 or not scopes or any(symbol.startswith("Scer\\GAL80[") for symbol in tools):
            return None, "", "RNAi without one whole-body or pan-neuronal driver"
        return genes.pop(), "FlyBase-RNAi", scopes[0]
    if tools or any("UAS" in symbol or "[+t" in symbol for symbol in rest):
        return None, "", "a transgene or a driven construct of the gene (overexpression or rescue)"
    return genes.pop(), "FlyBase-classical", "classical"


def build_flybase(path: Path, orthologs: dict, node_of_symbol: dict[str, str], crosswalk: list[SpeciesCrosswalkRow], drivers, validities) -> tuple[list[dict], dict]:
    human_of_symbol = by_symbol(orthologs)
    counts: dict = defaultdict(int)
    rows = []
    with gzip.open(path, "rt", encoding="utf-8") as handle:
        for line in handle:
            if line.startswith("#") or not line.strip():
                continue
            fields = line.rstrip("\n").split("\t")
            genotype, allele_ids, name, term, qualifiers, reference = fields[0], fields[1], fields[2], fields[3], fields[4], fields[6] if len(fields) > 6 else ""
            mapped = matching(crosswalk, "flybase", term, "", 1, validities)
            if not mapped:
                continue
            counts["mapped_rows"] += 1
            gene, source, scope = fly_genotype(re.split(r"[ /]+", genotype.strip()), drivers)
            if gene is None:
                counts[f"not_read:{scope}"] += 1
                continue
            human = human_of_symbol.get(gene)
            if human is None or human not in node_of_symbol:
                counts["no_one_to_one_human_gene_in_the_graph"] += 1
                continue
            direction = "loss_of_function" if source == "FlyBase-RNAi" else "untyped"
            for crosswalk_row in mapped:
                row = gene_row(source, human, node_of_symbol[human], direction, crosswalk_row.target_symptom, crosswalk_row.relation, crosswalk_row.validity, 1, gene)
                row.update({"report_id": f"{source}|{human}|{genotype}|{term}|{crosswalk_row.target_symptom}|{crosswalk_row.relation}|{reference}",
                            "source_record_id": allele_ids, "source_record_label": genotype, "source_term_id": term, "source_term_label": name,
                            "evidence_code": "flybase_call", "references": reference, "model_description": f"fly {genotype} ({scope}; {qualifiers or 'no qualifier'}); {name}; human orthologue {human} by Alliance one to one",
                            "limitations": f"cross-species: a fly {name} read as {crosswalk_row.target_symptom} with {crosswalk_row.validity} validity ({crosswalk_row.note})"})
                rows.append(row)
                counts["calls"] += 1
    return rows, dict(counts)


# --- zebrafish: ZFIN ------------------------------------------------------------------------------------------------


def read_fish_components(path: Path) -> dict[str, dict]:
    """fish id -> its genes, alleles, constructs and sequence-targeting reagents. The file's fifth column holds a genomic
    feature (ZDB-ALT) or a reagent (ZDB-MRPHLNO, ZDB-CRISPR, ZDB-TALEN); the seventh a construct."""
    fish: dict[str, dict] = defaultdict(lambda: {"genes": {}, "alleles": set(), "constructs": set(), "reagents": set()})
    with open(path, encoding="utf-8") as handle:
        for line in handle:
            fields = line.rstrip("\n").split("\t")
            if len(fields) < 12:
                continue
            record = fish[fields[0]]
            if fields[2]:
                record["genes"][fields[2]] = fields[3]
            if fields[4].startswith(("ZDB-MRPHLNO", "ZDB-CRISPR", "ZDB-TALEN")):
                record["reagents"].add(fields[4])
            elif fields[4]:
                record["alleles"].add(fields[4])
            if fields[6]:
                record["constructs"].add(fields[6])
            if fields[8]:
                record["reagents"].add(fields[8])
    return dict(fish)


def build_zfin(phenotype_path: Path, components: dict[str, dict], orthologs: dict, node_of_symbol: dict[str, str], crosswalk: list[SpeciesCrosswalkRow],
               validities) -> tuple[list[dict], dict]:
    counts: dict = defaultdict(int)
    rows = []
    calls_of: dict[tuple[str, str], bool] = {}
    candidates = []
    with open(phenotype_path, encoding="utf-8") as handle:
        for line in handle:
            fields = line.rstrip("\n").split("\t")
            if len(fields) < 23 or fields[6]:  # a post-composed subterm narrows the entity; not read
                continue
            term, quality, tag, environment = fields[10], fields[12], fields[14], fields[22]
            value = 0 if tag == "normal" else 1
            mapped = matching(crosswalk, "zfin", term, quality, value, validities)
            if not mapped:
                continue
            counts["mapped_rows"] += 1
            if environment not in ZFIN_ENVIRONMENTS_READ:
                counts["not_read:a treated or non-standard environment"] += 1
                continue
            record = components.get(fields[0])
            if record is None or len(record["genes"]) != 1 or record["constructs"]:
                counts["not_read:not one gene, or a transgenic construct"] += 1
                continue
            reagents = {reagent for reagent in record["reagents"] if reagent.startswith("ZDB-MRPHLNO")}
            if record["reagents"] - reagents or (record["alleles"] and reagents):
                counts["not_read:a CRISPR or TALEN reagent, or an allele with a morpholino"] += 1
                continue
            source = "ZFIN-morpholino" if reagents else "ZFIN-mutant"
            gene_id = next(iter(record["genes"]))
            ortholog = orthologs.get(f"ZFIN:{gene_id}")
            if ortholog is None or ortholog[2] not in node_of_symbol:
                counts["no_one_to_one_human_gene_in_the_graph"] += 1
                continue
            for crosswalk_row in mapped:
                key = (fields[0], crosswalk_row.target_symptom)
                if value == 1:
                    calls_of[key] = True
                candidates.append((fields, record, source, ortholog, crosswalk_row, value))
    for fields, record, source, ortholog, crosswalk_row, value in candidates:
        if value == 0 and calls_of.get((fields[0], crosswalk_row.target_symptom)):
            counts["absences_blocked_by_a_call_of_the_fish"] += 1
            continue
        human = ortholog[2]
        direction = "loss_of_function" if source == "ZFIN-morpholino" else "untyped"
        row = gene_row(source, human, node_of_symbol[human], direction, crosswalk_row.target_symptom, crosswalk_row.relation if value else "induces", crosswalk_row.validity, value, ortholog[0])
        row.update({"report_id": f"{source}|{human}|{fields[0]}|{fields[10]}|{fields[12]}|{crosswalk_row.target_symptom}|{crosswalk_row.relation if value else 'absent'}|{fields[21]}",
                    "source_record_id": fields[0], "source_record_label": html.unescape(re.sub(r"<[^>]+>", "", fields[1])), "source_term_id": f"{fields[10]}|{fields[12]}",
                    "source_term_label": f"{fields[11]}, {fields[13]} ({fields[14]})", "evidence_code": "zfin_call" if value else "zfin_normal", "references": fields[21],
                    "model_description": f"zebrafish {fields[1]} at {fields[3]} to {fields[5]}; {fields[11]} {fields[13]}, {fields[14]}; human orthologue {human} by Alliance one to one",
                    "limitations": f"cross-species: a zebrafish {fields[11]} ({fields[13]}) read as {crosswalk_row.target_symptom} with {crosswalk_row.validity} validity ({crosswalk_row.note})"})
        rows.append(row)
        counts["absences" if value == 0 else "calls"] += 1
    return rows, dict(counts)


# --- worm: WormBase -------------------------------------------------------------------------------------------------


def build_wormbase(path: Path, orthologs: dict, node_of_symbol: dict[str, str], crosswalk: list[SpeciesCrosswalkRow], validities) -> tuple[list[dict], dict]:
    counts: dict = defaultdict(int)
    raw = []
    called: set[tuple[str, str, str]] = set()
    with open(path, encoding="utf-8") as handle:
        for line in handle:
            if line.startswith("!") or not line.strip():
                continue
            fields = line.rstrip("\n").split("\t")
            negated = fields[3] == "NOT"
            mapped = matching(crosswalk, "wormbase", fields[4], "", 0 if negated else 1, validities)
            if not mapped:
                continue
            counts["mapped_rows"] += 1
            evidence_with = fields[7]
            source = "WormBase-RNAi" if "WBRNAi" in evidence_with else ("WormBase-variation" if "WBVar" in evidence_with else None)
            if source is None:
                counts["not_read:neither a variation nor an RNAi experiment"] += 1
                continue
            ortholog = orthologs.get(f"WB:{fields[1]}")
            if ortholog is None or ortholog[2] not in node_of_symbol:
                counts["no_one_to_one_human_gene_in_the_graph"] += 1
                continue
            for crosswalk_row in mapped:
                if not negated:
                    called.add((source, ortholog[2], crosswalk_row.target_symptom))
                raw.append((fields, source, ortholog, crosswalk_row, negated))
    rows = []
    for fields, source, ortholog, crosswalk_row, negated in raw:
        human = ortholog[2]
        if negated and (source, human, crosswalk_row.target_symptom) in called:
            counts["absences_blocked_by_a_call_of_the_gene"] += 1
            continue
        direction = "loss_of_function" if source == "WormBase-RNAi" else "untyped"
        row = gene_row(source, human, node_of_symbol[human], direction, crosswalk_row.target_symptom, "induces" if negated else crosswalk_row.relation, crosswalk_row.validity,
                       0 if negated else 1, ortholog[0])
        row.update({"report_id": f"{source}|{human}|{fields[1]}|{fields[4]}|{fields[7]}|{crosswalk_row.target_symptom}|{'absent' if negated else crosswalk_row.relation}|{fields[5]}",
                    "source_record_id": fields[7], "source_record_label": fields[2], "source_term_id": fields[4], "source_term_label": ("NOT " if negated else "") + fields[4],
                    "evidence_code": "wormbase_not" if negated else "wormbase_call", "references": fields[5],
                    "model_description": f"worm {fields[2]} ({fields[7]}); {'NOT ' if negated else ''}{fields[4]}; human orthologue {human} by Alliance one to one",
                    "limitations": f"cross-species: a worm readout read as {crosswalk_row.target_symptom} with {crosswalk_row.validity} validity ({crosswalk_row.note})"})
        rows.append(row)
        counts["absences" if negated else "calls"] += 1
    return rows, dict(counts)


# --- fly: the Fenckova et al. 2019 habituation screen ---------------------------------------------------------------

FENCKOVA_S2_COLUMNS = ("human_gene", "cg_number", "flybase_id", "fly_symbol", "vdrc_order_number", "lethality", "validated_by_repeat", "n_jumpers",
                       "percent_jumpers", "ttc_fold_change", "lm_estimate", "lm_standard_error", "p_value", "p_value_fdr", "fatigue", "category_line",
                       "gene_not_affected", "gene_habituation_deficient", "gene_non_performers", "gene_premature_habituation")
FENCKOVA_GENE_CATEGORIES = ("not affected", "habituation deficient", "non-performers", "premature habituation")


def read_fenckova_workbook(path: Path) -> tuple[list[dict], dict[str, str]]:
    """(Table S2 as one dict per RNAi line in FENCKOVA_S2_COLUMNS, Table S3 as human gene -> its gene-level category).
    S2 has two header rows; S3 one column per category ("non-peformers" sic in the file)."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = []
    for values in list(workbook["S2"].iter_rows(values_only=True))[2:]:
        if not values or not values[0]:
            continue
        lines.append(dict(zip(FENCKOVA_S2_COLUMNS, values[: len(FENCKOVA_S2_COLUMNS)])))
    s3 = list(workbook["S3"].iter_rows(values_only=True))
    category_of: dict[str, str] = {}
    names = [str(name).strip().lower().replace("non-peformers", "non-performers") for name in s3[0]]
    for values in s3[1:]:
        for name, gene in zip(names, values):
            if gene:
                category_of.setdefault(str(gene).strip(), name)
    return lines, category_of


def gene_category_from_s2(line: dict) -> str | None:
    flags = [category for category, column in zip(FENCKOVA_GENE_CATEGORIES, ("gene_not_affected", "gene_habituation_deficient", "gene_non_performers", "gene_premature_habituation"))
             if str(line.get(column) or "0").strip() in ("1", "1.0")]
    return flags[0] if len(flags) == 1 else None


def current_hgnc_symbols(hgnc_path: Path) -> dict[str, str]:
    """A symbol -> the current HGNC symbol: itself when current, else the one current gene that lists it as a previous
    symbol (a previous symbol two genes list is left out). The 2019 table names some genes by symbols HGNC has since
    changed (B3GALTL is now B3GLCT, KIAA1279 KIFBP)."""
    current, previous = set(), defaultdict(set)
    with open(hgnc_path, encoding="utf-8") as handle:
        for row in csv.DictReader(handle, delimiter="\t"):
            current.add(row["symbol"])
            for old in str(row.get("prev_symbol") or "").split("|"):
                if old.strip():
                    previous[old.strip()].add(row["symbol"])
    mapping = {symbol: symbol for symbol in current}
    for old, symbols in previous.items():
        if old not in mapping and len(symbols) == 1:
            mapping[old] = next(iter(symbols))
    return mapping


def build_fenckova(lines: list[dict], category_of: dict[str, str], fly_orthologs: dict, node_of_symbol: dict[str, str], crosswalk: list[SpeciesCrosswalkRow],
                   validities, current_symbol: dict[str, str] | None = None) -> tuple[list[dict], dict]:
    """One report per human gene from its gene-level category (Table S3, which S2's gene-level columns repeat): habituation
    deficient a call and not affected an absence of the crosswalk's symptom; non-performers and premature habituation
    neither; a gene whose lines are all lethal is not read. Each row names its lines with their FDR-adjusted p. The
    table's human symbols are brought to HGNC's current ones (current_symbol), and the human-to-fly pairs are checked
    against Alliance one to one, the mismatches counted and listed."""
    counts: dict = defaultdict(int)
    human_of_fly = {identifier.split(":", 1)[1]: human for identifier, (_, _, human) in fly_orthologs.items()}
    lines_of: dict[str, list[dict]] = defaultdict(list)
    mismatches = []
    renamed = {}
    for line in lines:
        counts["lines"] += 1
        table_symbol = str(line["human_gene"]).strip()
        human = (current_symbol or {}).get(table_symbol, table_symbol)
        if human != table_symbol:
            renamed[table_symbol] = human
        lines_of[human].append(line)
        alliance = human_of_fly.get(str(line["flybase_id"]).strip())
        if alliance != human:
            mismatches.append({"human_gene": human, "flybase_id": line["flybase_id"], "fly_symbol": line["fly_symbol"], "alliance_one_to_one_human": alliance})
    rows = []
    for human, gene_lines in sorted(lines_of.items()):
        counts["genes"] += 1
        table_symbol = str(gene_lines[0]["human_gene"]).strip()
        category = category_of.get(table_symbol) or gene_category_from_s2(gene_lines[0])
        s2_category = gene_category_from_s2(gene_lines[0])
        if s2_category and category and s2_category != category:
            counts["genes_whose_S2_and_S3_categories_differ"] += 1
        if all(str(line["lethality"]).strip().upper() == "YES" for line in gene_lines):
            counts["not_read:lethal"] += 1
            continue
        value = {"habituation deficient": 1, "not affected": 0}.get(category or "")
        if value is None:
            counts[f"not_read:{category}"] += 1
            continue
        if human not in node_of_symbol:
            counts["human_gene_not_in_graph"] += 1
            continue
        for crosswalk_row in matching(crosswalk, "fenckova", category, "", value, validities):
            description = "; ".join(f"VDRC {line['vdrc_order_number']} ({line['category_line']}, padj {line['p_value_fdr']}, N {line['n_jumpers']})" for line in gene_lines)
            row = gene_row("Fenckova-2019-RNAi", human, node_of_symbol[human], "loss_of_function", crosswalk_row.target_symptom, "induces", crosswalk_row.validity, value,
                           str(gene_lines[0]["flybase_id"]))
            row.update({"report_id": f"Fenckova-2019-RNAi|{human}|{crosswalk_row.target_symptom}|{category}", "source_record_id": str(gene_lines[0]["flybase_id"]),
                        "source_record_label": str(gene_lines[0]["fly_symbol"]), "source_term_id": category, "source_term_label": f"light-off jump habituation: {category}",
                        "evidence_code": "fenckova_habituation_deficient" if value else "fenckova_not_affected", "references": "PMID:31272685",
                        "model_description": f"fly pan-neuronal RNAi (elav-Gal4, UAS-Dicer-2) of {gene_lines[0]['fly_symbol']}: {description}",
                        "limitations": f"cross-species: habituation of the light-off jump read as {crosswalk_row.target_symptom} with {crosswalk_row.validity} validity ({crosswalk_row.note})"})
            rows.append(row)
            counts["absences" if value == 0 else "calls"] += 1
    unique = {(entry["human_gene"], entry["flybase_id"]): entry for entry in mismatches}
    counts["human_fly_pairs"] = len({(str(line["human_gene"]).strip(), str(line["flybase_id"]).strip()) for line in lines})
    counts["human_fly_pairs_without_an_Alliance_one_to_one_orthologue"] = sum(entry["alliance_one_to_one_human"] is None for entry in unique.values())
    counts["human_fly_pairs_one_to_one_with_another_human_gene"] = sum(entry["alliance_one_to_one_human"] is not None for entry in unique.values())
    counts["human_symbols_renamed_by_HGNC"] = len(renamed)
    return rows, {**dict(counts), "renamed": renamed, "mismatches": sorted(unique.values(), key=lambda entry: (entry["human_gene"], str(entry["flybase_id"])))}
